import os, re, io
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OLD = os.path.join(HERE, 'FrameFlow_Detailed_Project_Report (1).docx')
OUT = os.path.join(HERE, 'FrameFlow_Project_Report.docx')
EVID = os.path.join(ROOT, 'docs', 'evidence')

# ---------------------------------------------------------------- source content
src = Document(OLD)
imgs = []
blocks = []  # (kind, payload)
for el in src.element.body:
    if el.tag == qn('w:p'):
        p = docx.text.paragraph.Paragraph(el, src)
        blips = el.findall('.//' + qn('a:blip'))
        if blips:
            rid = blips[0].get(qn('r:embed'))
            blocks.append(('img', src.part.related_parts[rid].blob))
            continue
        t = p.text.strip()
        if not t:
            continue
        s = p.style.name
        if s == 'Heading 1': blocks.append(('h1', t))
        elif s == 'Heading 2': blocks.append(('h2', t))
        elif s == 'List Bullet': blocks.append(('bullet', t))
        elif s == 'List Number': blocks.append(('num', t))
        else: blocks.append(('p', t))
    elif el.tag == qn('w:tbl'):
        t = docx.table.Table(el, src)
        rows = [[c.text.strip().replace('\n', ' ') for c in r.cells] for r in t.rows]
        blocks.append(('table', rows))


def section_blocks(prefix, level):
    """Blocks of the old section whose heading starts with prefix (e.g. '5.1' or '5.')."""
    out, on = [], False
    for k, v in blocks:
        if k in ('h1', 'h2'):
            if on and (k == 'h1' or level == 2):
                break
            if v.startswith(prefix + ' ') or v.startswith(prefix + '.') and level == 1 and v.startswith(prefix + '. '):
                if (level == 1 and k == 'h1') or (level == 2 and k == 'h2'):
                    on = True
                    continue
            continue_flag = False
        elif on:
            out.append((k, v))
    return out


def h1_intro(num):
    out, on = [], False
    for k, v in blocks:
        if k == 'h1':
            if on: break
            on = v.startswith(f'{num}. ')
        elif k == 'h2':
            if on: break
        elif on:
            out.append((k, v))
    return out


def h2_body(num):
    out, on = [], False
    for k, v in blocks:
        if k in ('h1', 'h2'):
            if on: break
            on = (k == 'h2' and v.startswith(f'{num} '))
        elif on:
            out.append((k, v))
    return out

# ---------------------------------------------------------------- document setup
doc = Document()
FONT = 'Times New Roman'


def set_font(style, size, bold=False, italic=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), FONT)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]


st = doc.styles
n = st['Normal']; set_font(n, 12)
n.paragraph_format.line_spacing = 1.5
n.paragraph_format.space_after = Pt(6)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size, before, after in (('Heading 1', 16, 18, 12), ('Heading 2', 14, 12, 6), ('Heading 3', 13, 10, 6)):
    s = st[name]; set_font(s, size, bold=True)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.keep_with_next = True
for name in ('List Bullet', 'List Number'):
    s = st[name]; set_font(s, 12)
    s.paragraph_format.left_indent = Inches(0.5)
    s.paragraph_format.first_line_indent = Inches(-0.25)
    s.paragraph_format.line_spacing = 1.5
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(sec, m, Inches(1))


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)
    tcPr.append(sh)


def cell_margins(table, tb=120, lr=160):
    tblPr = table._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for side, v in (('top', tb), ('left', lr), ('bottom', tb), ('right', lr)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblPr.append(m)


def para(text='', style=None, align=None, bold=False, italic=False, size=None, after=None, before=None, spacing=None, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    if spacing is not None: p.paragraph_format.line_spacing = spacing
    if keep: p.paragraph_format.keep_with_next = True
    return p


C = WD_ALIGN_PARAGRAPH.CENTER
L = WD_ALIGN_PARAGRAPH.LEFT


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def front_title(text):
    para(text, align=C, bold=True, size=16, after=12, spacing=1.15, keep=True)


list_tables, list_figures = [], []
tcount = {}; fcount = {}


def caption(kind, ch, title):
    d = tcount if kind == 'Table' else fcount
    d[ch] = d.get(ch, 0) + 1
    label = f'{kind} {ch}.{d[ch]}: {title}'
    (list_tables if kind == 'Table' else list_figures).append(label)
    para(label, align=C, italic=True, size=11, after=10, spacing=1.15)


def add_table(rows, ch, title, widths=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(t)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ''
            p = c.paragraphs[0]
            p.alignment = L
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val); r.font.size = Pt(12); r.bold = (i == 0)
            if i == 0: shade(c, 'D5E8F0')
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Inches(w)
    # repeat header
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
    para('', after=0, spacing=1.0)
    caption('Table', ch, title)


def add_image(blob_or_path, ch, title, width=5.8):
    p = para(align=C, after=4, spacing=1.0, keep=True)
    src_ = io.BytesIO(blob_or_path) if isinstance(blob_or_path, bytes) else blob_or_path
    p.add_run().add_picture(src_, width=Inches(width))
    caption('Figure', ch, title)


def callout(rows):
    # old callout is a 1-row table: [title, text]
    title, text = rows[0][0], rows[0][-1]
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(t, 120, 200)
    c = t.cell(0, 0); shade(c, 'EEF5F9')
    p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.15; p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(title + ': '); r.bold = True
    p.add_run(text if text != title else '')
    para('', after=0, spacing=1.0)


TABLE_TITLES = {
    'Term': 'Abbreviations and Key Terms',
    'ID': 'Functional Requirements',
    'Element': 'Inputs and Outputs of the Interpolation Model',
    'Module': 'Repository Modules and Their Purpose',
    'Endpoint': 'Backend API Endpoints',
    'Component': 'Verified Development Environment',
    'Metric': None,
    'Method': 'Synthetic Demo Comparison of Interpolation Methods',
    'Measurement': 'Inference and Browser Performance',
    'Risk / limitation': 'Limitations and Risk Areas',
}
METRIC_TITLES = {'MAE': 'Evaluation Metrics and Interpretation'}


def emit(items, ch, h3=None, fig_titles=None):
    fig_titles = list(fig_titles or [])
    for k, v in items:
        if k == 'p':
            if re.match(r'^Figure \d+\.', v):
                continue
            para(v)
        elif k == 'bullet': para(v, style='List Bullet')
        elif k == 'num': para(v, style='List Number')
        elif k == 'img': add_image(v, ch, fig_titles.pop(0) if fig_titles else 'Illustration')
        elif k == 'table':
            if len(v) == 1:
                callout(v)
            else:
                first = v[0][0]
                if first == 'Metric' and v[0][1] == 'Interpretation': title = 'Evaluation Metrics and Interpretation'
                elif first == 'Metric' and v[0][1] == 'RIFE': title = 'GOES-16 RGB Ten-Sequence Validation Results'
                elif first == 'Metric': title = 'GOES-16 ABI Band 13 Thermal Evaluation Results'
                elif first == 'Method' and v[0][1] == 'Path': title = 'API Summary'
                elif first == 'Method' and len(v[0]) == 3: title = 'API Summary'
                else: title = TABLE_TITLES.get(first, 'Summary')
                add_table(v, ch, title)


def h1(text, new_page=True):
    if new_page: page_break_before_next()
    return doc.add_heading(text, level=1)


_pb = {'flag': False}


def page_break_before_next():
    _pb['flag'] = True


_orig_add_heading = doc.add_heading


def add_heading(text, level=1):
    h = _orig_add_heading(text, level=level)
    if level == 1 and _pb['flag']:
        h.paragraph_format.page_break_before = True
        _pb['flag'] = False
    return h


doc.add_heading = add_heading


def heading(text, level):
    return doc.add_heading(text, level=level)


def chapter(text):
    page_break_before_next(); return heading(text, 1)


# ---------------------------------------------------------------- headers / footers / page numbers
def add_page_field(paragraph):
    r = paragraph.add_run(); r.font.size = Pt(11); r.font.name = FONT
    for typ, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if typ:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), typ); r._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; r._r.append(it)


def set_page_numbering(section, fmt, start):
    sectPr = section._sectPr
    pg = sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType'); sectPr.append(pg)
    pg.set(qn('w:fmt'), fmt); pg.set(qn('w:start'), str(start))


def footer_with_number(section):
    section.footer.is_linked_to_previous = False
    f = section.footer
    for p in f.paragraphs: p.text = ''
    p = f.paragraphs[0]; p.alignment = C
    add_page_field(p)


def footer_blank(section):
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs: p.text = ''


# ================================================================= TITLE PAGE
para('Woxsen University', align=C, bold=True, size=22, before=30, after=2, spacing=1.15)
para('School of Technology', align=C, bold=True, size=16, after=50, spacing=1.15)
para('A', align=C, size=14, after=4, spacing=1.15)
para('PROJECT REPORT', align=C, bold=True, size=24, after=4, spacing=1.15)
para('on', align=C, size=14, after=14, spacing=1.15)
para('FRAMEFLOW: AI/ML OPTICAL-FLOW BASED TEMPORAL INTERPOLATION', align=C, bold=True, size=18, after=2, spacing=1.15)
para('OF GEOSTATIONARY SATELLITE IMAGERY USING RIFE HDv3', align=C, bold=True, size=18, after=8, spacing=1.15)
para('Problem Statement 12 - "Fill in the Frames Seamlessly"', align=C, italic=True, size=13, after=36, spacing=1.15)
para('Submitted in partial fulfillment of the requirements for the degree of', align=C, size=13, after=8, spacing=1.15)
para('B. Tech. in Artificial Intelligence and Machine Learning', align=C, bold=True, size=14, after=40, spacing=1.15)
para('Submitted by:', align=L, bold=True, size=13, after=2, spacing=1.15)
para('Atharva Kanojia and Team', align=L, size=13, after=0, spacing=1.15)
para('[Roll Number]', align=L, size=13, after=14, spacing=1.15)
para('Under the guidance of:', align=L, bold=True, size=13, after=2, spacing=1.15)
para('[Faculty Name]', align=L, size=13, after=14, spacing=1.15)
para('Academic Year 2026  |  Batch 2025 - 2029', align=C, size=12, before=20, spacing=1.15)

# ---- section 2: prelim pages (roman)
s2 = doc.add_section(WD_SECTION.NEW_PAGE)
set_page_numbering(s2, 'lowerRoman', 1)
footer_with_number(s2)
footer_blank(doc.sections[0])

# CERTIFICATE
front_title('CERTIFICATE')
para('This is to certify that the project report entitled "FrameFlow: AI/ML Optical-Flow Based Temporal Interpolation of Geostationary Satellite Imagery using RIFE HDv3" submitted by Atharva Kanojia and Team ([Roll Number]) in partial fulfillment of the requirements for the award of the degree of B.Tech. in Artificial Intelligence and Machine Learning from Woxsen University is a bonafide record of work carried out by the student(s) under my supervision and guidance.')
para('The work embodied in this project report has been carried out by the candidate(s) and has not been submitted elsewhere for a degree.')
para('', after=40)
para('Signature of Mentor', bold=True, align=L, spacing=1.15, after=2)
para('Name: [Mentor Name]', align=L, spacing=1.15, after=2)
para('Designation: [Designation]', align=L, spacing=1.15, after=2)
para('Date:', align=L, spacing=1.15)

# DECLARATION
page_break()
front_title('DECLARATION of the candidate')
para('I hereby declare that the project work entitled "FrameFlow: AI/ML Optical-Flow Based Temporal Interpolation of Geostationary Satellite Imagery using RIFE HDv3" submitted to the School of Technology, Woxsen University, in partial fulfillment of the requirements for the award of the degree of B.Tech. in Artificial Intelligence and Machine Learning is my original work and has been carried out under the guidance of [Guide Name].')
para('I further declare that the work reported in this project has not been submitted and will not be submitted, either in part or in full, for the award of any other degree or diploma in this institute or any other institute or university.')
for _ in range(2):
    para('', after=14)
    para('Signature of Student', bold=True, align=L, spacing=1.15, after=2)
    para('Name: [Your Name]', align=L, spacing=1.15, after=2)
    para('Roll Number: [Your Roll Number]', align=L, spacing=1.15, after=2)
para('Date:', align=L, spacing=1.15, before=10)
para('Note: Please add more signature blocks if applicable.', italic=True, size=11, align=L, spacing=1.15)

# ACKNOWLEDGMENT
page_break()
front_title('ACKNOWLEDGMENT')
para('I would like to express my sincere gratitude to all those who have contributed to the successful completion of this project on temporal interpolation of geostationary satellite imagery using the FrameFlow system.')
para('First and foremost, I extend my heartfelt thanks to my project guide, [Guide Name], [Designation], for their invaluable guidance, continuous support, and constructive feedback throughout the duration of this project. Their expertise in machine learning and computer vision has been instrumental in shaping this work.')
para('I am grateful to [Head of Department Name], Head of the Department of [Department Name], for providing the necessary facilities and resources required for this project.')
para('I would also like to thank the Indian Space Research Organisation (ISRO) for framing Problem Statement 12 of the Bharatiya Antariksh Hackathon 2026, the NOAA GOES-R programme for making GOES-16 data publicly available, and the authors of RIFE (Huang et al.) for releasing their frame interpolation model, which formed the foundation of this work.')
para('My sincere thanks to my peers and colleagues who provided valuable insights and suggestions during various phases of this project.')
para('Finally, I am deeply grateful to my family for their unwavering support and encouragement throughout my academic journey.')
para('Atharva Kanojia and Team', bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, before=12)

# ABSTRACT
page_break()
front_title('ABSTRACT')
for t in (
    'Temporal resolution is a recurring limitation in satellite-based weather monitoring. When a geostationary platform produces an image every fixed interval, fast cloud and storm evolution can occur between observations. This project, FrameFlow, studies whether a learned frame interpolation model can reconstruct useful intermediate satellite imagery without requiring another satellite observation. It was developed against ISRO Bharatiya Antariksh Hackathon 2026 Problem Statement 12.',
    'The primary objective of this work is to develop and evaluate an optical-flow based pipeline that takes two consecutive satellite frames and generates the missing intermediate frame. The project uses the pretrained RIFE HDv3 (IFNet) model as the interpolation engine, adapts the input path for satellite imagery including single-channel thermal brightness temperature, and evaluates the generated frame against a withheld ground-truth frame using MAE, RMSE, PSNR and SSIM, compared against a linear interpolation baseline.',
    'The methodology encompasses image preprocessing and normalization, RIFE inference, a confidence proxy derived from forward/backward warping residuals, quantitative evaluation, a FastAPI backend and a React/Vite dashboard with browser-level end-to-end verification.',
    'Experimental results on GOES-16 Hurricane Ian data show that, across ten non-overlapping RGB temporal samples, RIFE obtained a mean MAE of 8.598 compared with 17.747 for linear interpolation, a mean PSNR of 24.839 dB versus 19.533 dB, and a mean SSIM of 0.8109 versus 0.6809. In the ABI Band 13 thermal experiment, the reconstructed brightness-temperature frame achieved 0.358 K MAE, 0.766 K RMSE, 45.836 dB PSNR and 0.9925 SSIM, all better than the linear baseline. These values demonstrate prototype-level feasibility and not operational meteorological accuracy.',
    'This research contributes a complete, reproducible workflow from frame input to interpolation, quality measurement and visual inspection, and identifies the next steps required for real INSAT deployment, broader cross-event validation and satellite-specific fine-tuning.',
):
    para(t)
para('Keywords: Satellite Imagery, Video Frame Interpolation, Optical Flow, RIFE, GOES-16, Brightness Temperature, Confidence Estimation, FastAPI, React', italic=True, align=L)

# TABLE OF CONTENTS (field)
page_break()
front_title('TABLE OF CONTENTS')
p = para(align=L, spacing=1.15)
r = p.add_run()
for typ in ('begin',):
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); fc.set(qn('w:dirty'), 'true'); r._r.append(fc)
it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'TOC \\o "1-3" \\h \\z \\u'; r._r.append(it)
fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'separate'); r._r.append(fc)
t = OxmlElement('w:t'); t.text = 'Right-click and choose "Update Field" to generate the table of contents.'; r._r.append(t)
fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'end'); r._r.append(fc)

# placeholders for lists (filled at end)
page_break()
front_title('LIST OF TABLES')
LT_ANCHOR = para('@@LT@@', align=L, spacing=1.15)
page_break()
front_title('LIST OF FIGURES')
LF_ANCHOR = para('@@LF@@', align=L, spacing=1.15)

# ---- section 3: main body (arabic)
s3 = doc.add_section(WD_SECTION.NEW_PAGE)
set_page_numbering(s3, 'decimal', 1)
footer_with_number(s3)

# ================================================================= CHAPTER 1
heading('1. INTRODUCTION', 1)
para('The temporal gaps between consecutive geostationary satellite observations represent a significant limitation in weather monitoring. With the increasing availability of high-cadence satellite datasets and advances in deep learning-based video frame interpolation, estimating the missing observation computationally has become increasingly feasible. This project focuses on developing a robust interpolation pipeline, FrameFlow, that generates an intermediate frame between two observed satellite frames and verifies it quantitatively.')
heading('1.1 Background', 2)
emit(h2_body('1.1'), 1, fig_titles=['Basic Temporal Interpolation Problem Addressed by FrameFlow'])
emit(h2_body('1.2'), 1)
heading('1.2 Motivation', 2)
para('The motivation for this project stems from several practical and research-oriented considerations:')
emit(h2_body('1.3'), 1)
heading('1.3 Problem Statement', 2)
emit(h2_body('2.1'), 1)
emit(h2_body('2.2'), 1)
heading('1.4 Objectives', 2)
para('The specific objectives of this project are:')
emit(h2_body('2.3'), 1)
heading('1.5 Dataset Overview', 2)
para('The project evaluates the system on two real data experiments drawn from GOES-16 observations of Hurricane Ian (27 September 2022), in addition to a synthetic end-to-end demo with the following characteristics:')
add_table([
    ['Attribute', 'RGB Validation', 'Band 13 Thermal Experiment'],
    ['Satellite', 'GOES-16 (GOES-East)', 'GOES-16 ABI Level-2 CMIP'],
    ['Product', 'GEOCOLOR sequence', 'Band 13 (10.3 micrometer clean IR)'],
    ['Event', 'Hurricane Ian, 27 Sep 2022', 'Hurricane Ian, 27 Sep 2022'],
    ['Samples', '10 non-overlapping triplets', '1 triplet (T0, T1, T2)'],
    ['Region / size', 'Full-frame sequence', '512 x 512 crop'],
    ['Units', 'RGB pixel values (0-255)', 'Brightness temperature (K)'],
], 1, 'Dataset Statistics', widths=[1.5, 2.3, 2.7])
heading('1.6 Scope and Limitations', 2)
emit(h2_body('2.4'), 1)
para('Limitations of the present scope include:')
for b in ('Real-data validation is centered on a single weather event (Hurricane Ian)',
          'The pretrained RIFE model has not been fine-tuned on a large multi-event satellite corpus',
          'Operational INSAT-3DS/3DR performance has not been established',
          'The confidence map is a consistency proxy rather than a calibrated uncertainty estimate'):
    para(b, style='List Bullet')
heading('1.7 Organization of the Report', 2)
para('The remainder of this report is organized as follows: Chapter 2 presents a review of related work in optical flow, frame interpolation and satellite-specific interpolation. Chapter 3 details the methodology, including system requirements, architecture, data preprocessing, the RIFE model, confidence estimation, software implementation and evaluation design. Chapter 4 presents the experimental results and discussion. Chapter 5 concludes the report with a summary of contributions, limitations and directions for future work.')

# ================================================================= CHAPTER 2
chapter('2. LITERATURE REVIEW / TECHNOLOGY REVIEW')
para('This chapter reviews existing research and technology in the domains of optical flow, video frame interpolation and satellite image temporal interpolation, with particular emphasis on the RIFE family of models used in this project.')
heading('2.1 Optical Flow and Frame Interpolation', 2)
emit(h2_body('3.1'), 2)
heading('2.2 RIFE / IFNet', 2)
emit(h2_body('3.2'), 2)
heading('2.3 Why RIFE Was Selected', 2)
emit(h2_body('3.3'), 2)
heading('2.4 Prior Satellite-Specific Work', 2)
emit(h2_body('3.4'), 2)
heading('2.5 Evaluation Metrics in Interpolation Studies', 2)
para('Standard full-reference metrics for image reconstruction include mean absolute error (MAE), root mean squared error (RMSE), peak signal-to-noise ratio (PSNR) and the structural similarity index (SSIM). For thermal satellite products, errors are additionally reported in physical units (Kelvin) so that the result retains meteorological meaning rather than only display-domain similarity.')
heading('2.6 Research Gap', 2)
para('While significant progress has been made in generic video frame interpolation, several areas warrant further investigation for satellite data: handling single-channel physical fields, providing a visible confidence indicator for generated pixels, and packaging the whole workflow with a measurable evaluation loop and a usable interface. This project addresses these gaps by integrating a pretrained RIFE model with satellite-aware preprocessing, a flow-consistency confidence proxy, and an API and dashboard.')

# ================================================================= CHAPTER 3
chapter('3. METHODOLOGY')
para('This chapter details the methodology employed in developing FrameFlow. It covers the system requirements, architecture, data preparation, model integration, confidence estimation, software implementation and the evaluation strategy.')
heading('3.1 System Overview', 2)
emit(h1_intro('5'), 3, fig_titles=['High-Level Processing Path Implemented by FrameFlow'])
heading('3.2 System Requirements and Design Goals', 2)
heading('3.2.1 Functional Requirements', 3)
emit(h2_body('4.1'), 3)
heading('3.2.2 Non-Functional Requirements', 3)
emit(h2_body('4.2'), 3)
heading('3.3 Proposed System Architecture', 2)
for i, name in enumerate(('Data Layer', 'Preprocessing Layer', 'Model Layer', 'Service Layer', 'Presentation Layer'), 1):
    heading(f'3.3.{i} {name}', 3)
    emit(h2_body(f'5.{i}'), 3)
heading('3.4 Data Sources and Preprocessing', 2)
for i, name in enumerate(('Intended Data Strategy', 'GOES-16 RGB Validation Dataset', 'ABI Band 13 Thermal Dataset', 'Normalization and Physical Units', 'Registration and Masking'), 1):
    heading(f'3.4.{i} {name}', 3)
    emit(h2_body(f'6.{i}'), 3, fig_titles=['Thermal Band 13 Processing Path Used in the Prototype Experiment'])
heading('3.5 RIFE Model and Interpolation Method', 2)
for i, name in enumerate(('RIFE HDv3 / IFNet', 'Inputs and Outputs', 'Linear Baseline', 'Why the RIFE Architecture Was Not Modified'), 1):
    heading(f'3.5.{i} {name}', 3)
    emit(h2_body(f'7.{i}'), 3)
heading('3.6 Confidence Estimation and Error Visualization', 2)
for i, name in enumerate(('Motivation', 'Computation'), 1):
    heading(f'3.6.{i} {name}', 3)
    emit(h2_body(f'8.{i}'), 3)
heading('3.7 Software Implementation', 2)
for i, name in enumerate(('Repository Organization', 'Backend API', 'Frontend Workflow', 'Development Environment'), 1):
    heading(f'3.7.{i} {name}', 3)
    emit(h2_body(f'9.{i}'), 3)
heading('3.8 Evaluation Methodology', 2)
for i, name in enumerate(('Evaluation Setup', 'Metrics', 'Baselines', 'Test Coverage'), 1):
    heading(f'3.8.{i} {name}', 3)
    emit(h2_body(f'10.{i}'), 3, fig_titles=['Main Validation Stages Recorded in the Project Documentation'])

# ================================================================= CHAPTER 4
chapter('4. RESULTS AND DISCUSSION')
para('This chapter presents the experimental results obtained from evaluating the FrameFlow pipeline. The results are analyzed in terms of reconstruction accuracy against the linear baseline, physical-unit error for thermal imagery, and inference and application performance.')
for i, name in enumerate(('Synthetic End-to-End Demo', 'GOES-16 RGB Multi-Sequence Validation', 'GOES-16 ABI Band 13 Thermal Evaluation', 'Inference and Browser Performance'), 1):
    heading(f'4.{i} {name}', 2)
    emit(h2_body(f'11.{i}'), 4)
heading('4.5 Dashboard Output and Visual Evidence', 2)
para('The React/Vite dashboard was exercised end to end in a Chromium browser. The following screenshots, captured from the working application, show the empty dashboard, the uploaded bracketing frames, the full result view, the metrics section and the invalid-input error path.')
ev = [('1_dashboard_empty.png', 'Dashboard in Its Initial State'),
      ('2_uploaded_frames.png', 'Uploaded Bracketing Frames'),
      ('3_results_full.png', 'Full Interpolation Result View'),
      ('5_metrics_section.png', 'Metrics Section of the Dashboard'),
      ('6_error_flow.png', 'Handling of Invalid Input')]
for fn, title in ev:
    pth = os.path.join(EVID, fn)
    if os.path.exists(pth):
        add_image(pth, 4, title, width=5.6)
heading('4.6 Discussion', 2)
for i, name in enumerate(('What Worked Well', 'Why the Baseline Still Matters', 'Where the Model Is Likely to Struggle'), 1):
    heading(f'4.6.{i} {name}', 3)
    emit(h2_body(f'12.{i}'), 4)
heading('4.7 Limitations and Risk Areas', 2)
emit(h1_intro('13'), 4)
heading('4.7.1 Research Integrity of the Current Claims', 3)
emit(h2_body('13.1'), 4)

# ================================================================= CHAPTER 5
chapter('5. CONCLUSION AND FUTURE WORK')
heading('5.1 Summary of Contributions', 2)
para('This project successfully developed and evaluated an optical-flow based temporal interpolation system for geostationary satellite imagery. The key contributions of this work include:')
for i, t in enumerate((
    'Integration of the pretrained RIFE HDv3 model with a satellite-aware preprocessing path for RGB and single-channel thermal imagery.',
    'Quantitative validation showing RIFE outperforming linear interpolation in all ten GOES-16 RGB samples (mean PSNR 24.839 dB vs 19.533 dB, mean SSIM 0.8109 vs 0.6809).',
    'A GOES-16 ABI Band 13 experiment reporting 0.358 K MAE and 0.9925 SSIM on the tested crop, evaluated in physical brightness-temperature units.',
    'A flow-consistency confidence proxy that exposes regions where the generated frame is less reliable.',
    'A complete software stack: FastAPI service, React/Vite dashboard, unit tests and a browser end-to-end test.'), 1):
    para(t, style='List Number')
emit([('p', x) for k, x in h1_intro('15')[:2]], 5)
heading('5.2 Limitations', 2)
para('While the project achieved its objectives, several limitations should be acknowledged:')
for b in ('Real-data validation is centered on one storm sequence (Hurricane Ian), which limits claims of cross-event generalization.',
          'The Band 13 experiment is based on a single 512 x 512 crop.',
          'The confidence map is a structural consistency indicator, not a calibrated probability of correctness.',
          'The documented results do not establish operational INSAT-3DS/3DR performance.'):
    para(b, style='List Bullet')
heading('5.3 Future Work', 2)
for i, name in enumerate(('Satellite-Specific Fine-Tuning', 'Cross-Event and Cross-Satellite Validation', 'INSAT-3DS/3DR Domain Adaptation', 'Better Uncertainty Estimation', 'Deployment Optimization'), 1):
    heading(f'5.3.{i} {name}', 3)
    emit(h2_body(f'14.{i}'), 5)
heading('5.4 Concluding Remarks', 2)
for k, x in h1_intro('15')[2:]:
    para(x)

# ================================================================= REFERENCES
page_break_before_next()
heading('REFERENCES', 1)
refs = [x for k, x in blocks if k == 'p' and re.match(r'^\[\d+\] ', x)]
for r_ in refs:
    p = para(r_, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.4)

# ================================================================= APPENDICES
page_break_before_next()
heading('APPENDIX - I', 1)
para('Screen shots of work done.', bold=True)
para('Screen shots of deployment tools of:')
for b in ('FrameFlow dashboard (React/Vite)', 'FastAPI backend (interactive API documentation)', 'GitHub repository'):
    para(b, style='List Bullet')
para('Dashboard screenshots are included in Section 4.5 (Figures 4.1 - 4.5). Additional repository and deployment screenshots may be inserted below.')
para('Appendix I-A: API Summary', bold=True, before=8)
emit([x for x in [('table', next(v for k, v in blocks if k == 'table' and v[0][0] == 'Method' and len(v[0]) == 3))]], 6)
para('Appendix I-B: Documentation Status Note', bold=True, before=8)
for k, x in blocks:
    if k == 'p' and x.startswith('The project folder contains phase-wise'):
        para(x)

page_break_before_next()
heading('APPENDIX - II', 1)
heading('FORMATTING GUIDELINES', 2)
para('This appendix summarizes the formatting guidelines followed in the preparation of this report, as per the Woxsen University project report template.')
add_table([['Element', 'Specification'],
           ['Paper Size', 'A4'], ['Margins', '1 inch (2.54 cm) on all sides'],
           ['Font Family', 'Times New Roman'], ['Body Text Font Size', '12 pt'],
           ['Line Spacing', '1.5'], ['Text Alignment', 'Justified'], ['Paragraph Spacing', '6 pt after paragraphs']],
          7, 'General Formatting Guidelines', widths=[2.5, 4.0])
para('Heading Styles', bold=True)
for b in ('Heading 1 (Chapter Titles): 16 pt, bold, black, 18 pt before, 12 pt after',
          'Heading 2 (Section Titles): 14 pt, bold, black, 12 pt before, 6 pt after',
          'Heading 3 (Subsection Titles): 13 pt, bold, black, 10 pt before, 6 pt after'):
    para(b, style='List Bullet')
para('Document Structure', bold=True)
for b in ('Title Page', 'Certificate', 'Declaration', 'Acknowledgment', 'Abstract', 'Table of Contents', 'List of Tables', 'List of Figures', 'Main Chapters (numbered)', 'References', 'Appendices'):
    para(b, style='List Number')

# ---------------------------------------------------------------- fill lists
def fill(anchor, items):
    anchor.text = ''
    first = True
    for it_ in items:
        if first:
            p = anchor; first = False
        else:
            new = OxmlElement('w:p'); p._p.addnext(new)
            p = docx.text.paragraph.Paragraph(new, anchor._parent)
            p.alignment = L; p.paragraph_format.line_spacing = 1.15
        p.add_run(it_)

# only chapter-style tables (exclude the 'Table 7.x' formatting table from lists? keep all)
fill(LT_ANCHOR, list_tables)
fill(LF_ANCHOR, list_figures)

# update fields on open
settings = doc.settings.element
uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true'); settings.append(uf)

doc.core_properties.title = 'FrameFlow Project Report'
doc.core_properties.author = 'Atharva Kanojia and Team'
doc.save(OUT)
print('saved', OUT, len(list_tables), 'tables', len(list_figures), 'figures')
